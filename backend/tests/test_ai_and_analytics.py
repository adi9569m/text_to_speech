import io
import docx
import pytest
from fastapi.testclient import TestClient
from backend.main import app

client = TestClient(app)


def test_ai_enhance_grammar():
    response = client.post(
        "/api/ai/enhance",
        json={"text": "um this is a test ???", "mode": "grammar"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["mode"] == "grammar"
    assert "This is a test?" in data["enhanced_text"]
    assert len(data["changes_applied"]) > 0


def test_ai_enhance_summarize():
    long_text = (
        "Artificial intelligence is transforming every industry across the globe today. "
        "It helps automate repetitive tasks and optimize complex business operations. "
        "Machine learning models can analyze vast amounts of data in seconds. "
        "Ultimately, these advancements allow humans to focus on creative and high-impact endeavors."
    )
    response = client.post(
        "/api/ai/enhance",
        json={"text": long_text, "mode": "summarize"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["mode"] == "summarize"
    assert len(data["enhanced_text"]) < len(long_text)


def test_ai_enhance_conversational():
    response = client.post(
        "/api/ai/enhance",
        json={"text": "We will not utilize obsolete tools.", "mode": "conversational"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "won't" in data["enhanced_text"] or "use" in data["enhanced_text"]


def test_ai_enhance_formal():
    response = client.post(
        "/api/ai/enhance",
        json={"text": "hey there! it's gonna work great", "mode": "formal"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "going to" in data["enhanced_text"] or "it is" in data["enhanced_text"]


def test_ai_enhance_bullet_to_script():
    bullet_text = "- First item to remember\n- Second crucial step\n- Final recommendation"
    response = client.post(
        "/api/ai/enhance",
        json={"text": bullet_text, "mode": "bullet_to_script"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "First" in data["enhanced_text"]


def test_ai_enhance_validation_errors():
    # Empty text
    res = client.post("/api/ai/enhance", json={"text": "", "mode": "grammar"})
    assert res.status_code in (400, 422)

    # Invalid mode
    res = client.post("/api/ai/enhance", json={"text": "Hello world", "mode": "nonexistent_mode"})
    assert res.status_code == 400


def test_analytics_endpoint():
    response = client.get("/api/analytics")
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "total_generations" in data
    assert "total_characters_synthesized" in data
    assert "top_voices" in data
    assert "top_languages" in data


def test_docx_document_extraction():
    doc = docx.Document()
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This is a paragraph from a Word document for speech synthesis.")
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)

    files = {
        "file": (
            "test_sample.docx",
            doc_buffer,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/extract-text", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "paragraph from a Word document" in data["text"]
    assert data["filename"] == "test_sample.docx"


def test_voice_sample_preview_endpoint():
    # JennyNeural
    response = client.get("/api/voices/en-US-JennyNeural/sample")
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "sample_audio_url" in data
    assert data["voice"] == "en-US-JennyNeural"

    # Nonexistent voice
    res_err = client.get("/api/voices/nonexistent-voice-xyz/sample")
    assert res_err.status_code == 404
